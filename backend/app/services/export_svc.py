"""Excel, CSV and Parquet output plus the Excel round trip (section 12)."""
from __future__ import annotations

import csv
import io
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.auth import Principal
from app.services import grid, layers
from app.services.grid import GridQuery

INK = "FF101826"
PAPER = "FFFAFAF7"
GREY = "FFE9E8E3"
FREEZE = "FFE3EBF5"
OVERRIDE = "FFF6E8CE"

THIN = Side(style="thin", color="FFD9D9D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def export_grid_xlsx(db: Session, principal: Principal, q: GridQuery) -> bytes:
    payload = grid.query(db, principal, q)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tahmin"

    header = ["Kırılım", "Katman"] + [c["label"] for c in payload["columns"]] + ["Yıl", "YTD", "YTG"]
    ws.append(header)
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFFFF", name="Calibri", size=10)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    editable_layers = {l["code"] for l in payload["layers"] if l["can_edit"]}
    flags = {c["bucket_id"]: c["flag"] for c in payload["columns"]}
    row_no = 2
    for row in payload["rows"]:
        for layer in payload["layers"]:
            code = layer["code"]
            ws.cell(row=row_no, column=1, value=row["label"]).border = BORDER
            ws.cell(row=row_no, column=2, value=layer["name"]).border = BORDER
            for i, col in enumerate(payload["columns"], start=3):
                v = row["layers"].get(code, {}).get(str(col["bucket_id"]))
                cell = ws.cell(row=row_no, column=i, value=v)
                cell.number_format = "#,##0"
                cell.border = BORDER
                locked = not (code in editable_layers and flags.get(col["bucket_id"]) in ("freeze", "open"))
                cell.protection = Protection(locked=locked)
                if flags.get(col["bucket_id"]) == "actual":
                    cell.fill = PatternFill("solid", fgColor=GREY)
                elif flags.get(col["bucket_id"]) == "fence":
                    cell.fill = PatternFill("solid", fgColor=GREY)
                elif flags.get(col["bucket_id"]) == "freeze":
                    cell.fill = PatternFill("solid", fgColor=FREEZE)
                if row["flags"].get(str(col["bucket_id"]), {}).get("override"):
                    cell.fill = PatternFill("solid", fgColor=OVERRIDE)
            sub = row["subtotals"].get(code, {})
            for j, k in enumerate(("year", "ytd", "ytg")):
                c = ws.cell(row=row_no, column=3 + len(payload["columns"]) + j, value=sub.get(k))
                c.number_format = "#,##0"
                c.border = BORDER
                c.protection = Protection(locked=True)
            row_no += 1

    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    for i in range(3, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.protection.sheet = True
    ws.protection.password = "forecast"

    meta = wb.create_sheet("Meta")
    meta["A1"] = "Bu sayfa geri yükleme için gereklidir, silmeyin."
    meta["A2"] = json.dumps(
        {
            "version_id": q.version_id,
            "level": q.level,
            "scope": q.scope,
            "grain": q.grain,
            "measure": q.measure,
            "layers": [l["code"] for l in payload["layers"]],
            "layer_names": {l["name"]: l["code"] for l in payload["layers"]},
            "columns": [{"bucket_id": c["bucket_id"], "label": c["label"], "flag": c["flag"]}
                        for c in payload["columns"]],
            "rows": [{"group": r["group"], "label": r["label"]} for r in payload["rows"]],
        },
        ensure_ascii=False,
    )
    meta.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_grid_xlsx(db: Session, principal: Principal, blob: bytes) -> dict:
    """Read an uploaded workbook and return the differences for confirmation."""
    wb = load_workbook(io.BytesIO(blob), data_only=True)
    if "Meta" not in wb.sheetnames:
        raise ValueError("Meta sayfası yok. Yalnızca platformdan indirilen şablon geri yüklenebilir.")
    meta = json.loads(wb["Meta"]["A2"].value)
    ws = wb["Tahmin"]

    columns = meta["columns"]
    layer_by_name = meta["layer_names"]
    editable_flags = {c["bucket_id"]: c["flag"] for c in columns}
    label_to_group = {r["label"]: r["group"] for r in meta["rows"]}

    q = GridQuery(
        version_id=meta["version_id"], scope=meta.get("scope") or {}, level=meta.get("level") or {},
        grain=meta.get("grain", "month"), layer_codes=meta["layers"], measure=meta.get("measure", "units"),
    )
    current = grid.query(db, principal, q)
    cur_index = {
        (r["group"], code): r["layers"].get(code, {})
        for r in current["rows"]
        for code in meta["layers"]
    }
    editable_layers = {l["code"] for l in current["layers"] if l["can_edit"]}

    diffs: list[dict] = []
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        group = label_to_group.get(str(row[0]))
        code = layer_by_name.get(str(row[1]))
        if group is None or code is None:
            continue
        for i, col in enumerate(columns, start=2):
            new = row[i] if i < len(row) else None
            if new is None or isinstance(new, str):
                continue
            if code not in editable_layers or editable_flags.get(col["bucket_id"]) not in ("freeze", "open"):
                skipped += 1
                continue
            old = cur_index.get((group, code), {}).get(str(col["bucket_id"]))
            if old is None or abs(float(new) - float(old)) > 0.5:
                diffs.append(
                    {
                        "group": group, "group_label": str(row[0]), "layer": code,
                        "bucket_id": col["bucket_id"], "bucket_label": col["label"],
                        "old": old, "new": round(float(new), 2),
                        "delta": round(float(new) - float(old or 0), 2),
                    }
                )
    return {
        "version_id": meta["version_id"], "level": meta.get("level") or {},
        "diffs": diffs, "skipped_locked_cells": skipped, "count": len(diffs),
    }


def apply_grid_diffs(
    db: Session, principal: Principal, payload: dict, reason_code_id: int, note: str = ""
) -> dict:
    """Apply confirmed differences as overrides."""
    from app.services.hierarchy import group_map
    from app.services.override_svc import OverrideRequest, apply_override

    level = payload.get("level") or {}
    gmap = group_map(db, level)
    applied, errors = 0, []
    for d in payload.get("diffs", []):
        keys = [sk for sk, g in gmap.items() if g == d["group"]]
        if not keys:
            errors.append(f"{d['group_label']}: seri bulunamadi")
            continue
        scope = _scope_from_group(d["group"])
        try:
            apply_override(
                db, principal,
                OverrideRequest(
                    version_id=payload["version_id"], layer_code=d["layer"], bucket_id=d["bucket_id"],
                    mode="absolute", value=float(d["new"]), reason_code_id=reason_code_id,
                    note=note or "Excel geri yükleme",
                    series_key=keys[0] if len(keys) == 1 else None,
                    scope=None if len(keys) == 1 else scope,
                    level=None if len(keys) == 1 else level,
                ),
            )
            applied += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{d['group_label']} {d['bucket_label']}: {exc}")
    return {"applied": applied, "errors": errors}


def _scope_from_group(group: str) -> dict:
    scope: dict = {}
    for part in group.split("|"):
        dim, nid = part.split("=", 1)
        scope[dim] = {"nodes": [int(nid)]}
    return scope


def rows_to_csv(columns: list[str], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns, delimiter=";", extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow(r)
    return buf.getvalue().encode("utf-8-sig")


def rows_to_parquet(columns: list[str], rows: list[dict]) -> bytes:
    import pandas as pd

    df = pd.DataFrame(rows, columns=columns)
    buf = io.BytesIO()
    try:
        df.to_parquet(buf, index=False)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(
            "Parquet yazimi için pyarrow gerekiyor. CSV veya Excel kullanın."
        ) from exc
    return buf.getvalue()


def table_to_xlsx(title: str, columns: list[dict], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:28] or "Rapor"
    ws.append([c["label"] for c in columns])
    for col in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.border = BORDER
    for r in rows:
        ws.append([r.get(c["key"]) for c in columns])
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c["label"]) + 6))
        if c.get("numeric"):
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=i).number_format = "#,##0.0"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
